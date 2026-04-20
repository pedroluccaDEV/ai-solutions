import io
import csv
import json
from typing import List, Dict, Any
from enum import Enum
from datetime import datetime

# =====================================================
# Bibliotecas opcionais (carregar só se necessário)
# =====================================================

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


# =====================================================
# Tipos de arquivo
# =====================================================

class FileType(Enum):
    TEXT = "text"
    TABLE = "table"
    PDF = "pdf"
    DOCX = "docx"
    CODE = "code"
    JSON = "json"
    IMAGE = "image"
    UNKNOWN = "unknown"


# =====================================================
# Engine principal
# =====================================================

class FileInterpretationEngine:
    """
    Motor bruto, rápido e escalável de interpretação de arquivos.
    Não depende de agente, planner ou LLM.
    """

    def __init__(self):
        self.supported_code_exts = {
            "py", "js", "ts", "java", "go", "cpp", "c", "rs",
            "php", "rb", "sh", "sql", "yaml", "yml"
        }

    # =====================================================
    # Public API
    # =====================================================

    async def interpret_files(self, files: List[Dict[str, Any]]) -> Dict[str, Any]:
        text_blocks = []
        table_blocks = []
        images = []

        for file in files:
            try:
                result = self._interpret_single_file(file)

                text_blocks.extend(result.get("text_blocks", []))
                table_blocks.extend(result.get("table_blocks", []))
                images.extend(result.get("images", []))

            except Exception as e:
                text_blocks.append({
                    "content": f"[Erro ao processar arquivo {file.get('filename')}: {str(e)}]",
                    "metadata": {
                        "type": "error",
                        "filename": file.get("filename")
                    }
                })

        combined_context = self._combine_text_blocks(text_blocks)

        return {
            "text_blocks": text_blocks,
            "table_blocks": table_blocks,
            "images": images,
            "combined_context": combined_context,
            "metadata": {
                "file_count": len(files),
                "text_block_count": len(text_blocks),
                "table_count": len(table_blocks),
                "image_count": len(images),
                "total_text_length": len(combined_context),
                "generated_at": datetime.utcnow().isoformat(),
                "pdf_parser_priority": ["pdfplumber", "pypdf2"]
            }
        }

    # =====================================================
    # Core
    # =====================================================

    def _interpret_single_file(self, file: Dict[str, Any]) -> Dict[str, Any]:
        filename = file.get("filename", "")
        content_type = file.get("content_type", "")

        file_type = self._detect_file_type(filename, content_type)

        if file_type == FileType.IMAGE:
            return self._interpret_image(file)

        if file_type == FileType.JSON:
            return self._interpret_json(file)

        if file_type == FileType.TABLE:
            return self._interpret_table(file)

        if file_type == FileType.PDF:
            return self._interpret_pdf(file)

        if file_type == FileType.DOCX:
            return self._interpret_docx(file)

        if file_type == FileType.CODE:
            return self._interpret_code(file)

        if file_type == FileType.TEXT:
            return self._interpret_text(file)

        return self._interpret_fallback(file)

    # =====================================================
    # Detecção de tipo
    # =====================================================

    def _detect_file_type(self, filename: str, content_type: str) -> FileType:
        ext = filename.lower().split(".")[-1]

        if content_type.startswith("image/"):
            return FileType.IMAGE

        if ext in {"csv", "tsv", "xlsx"}:
            return FileType.TABLE

        if ext == "pdf":
            return FileType.PDF

        if ext == "docx":
            return FileType.DOCX

        if ext == "json":
            return FileType.JSON

        if ext in self.supported_code_exts:
            return FileType.CODE

        if content_type.startswith("text/") or ext in {"txt", "md"}:
            return FileType.TEXT

        return FileType.UNKNOWN

    # =====================================================
    # Interpreters
    # =====================================================

    def _interpret_text(self, file):
        text = file["bytes"].decode("utf-8", errors="ignore")

        return {
            "text_blocks": [{
                "content": text,
                "metadata": {
                    "type": "text",
                    "filename": file["filename"]
                }
            }]
        }

    def _interpret_code(self, file):
        code = file["bytes"].decode("utf-8", errors="ignore")
        lang = file["filename"].split(".")[-1]

        return {
            "text_blocks": [{
                "content": f"```{lang}\n{code}\n```",
                "metadata": {
                    "type": "code",
                    "filename": file["filename"]
                }
            }]
        }

    def _interpret_json(self, file):
        raw = file["bytes"].decode("utf-8", errors="ignore")

        try:
            parsed = json.loads(raw)
            content = json.dumps(parsed, indent=2, ensure_ascii=False)
        except Exception:
            content = raw

        return {
            "text_blocks": [{
                "content": content,
                "metadata": {
                    "type": "json",
                    "filename": file["filename"]
                }
            }]
        }

    def _interpret_table(self, file):
        filename = file["filename"]
        ext = filename.split(".")[-1]

        tables = []
        text_blocks = []

        if ext in {"csv", "tsv"}:
            delimiter = "," if ext == "csv" else "\t"
            content = file["bytes"].decode("utf-8", errors="ignore").splitlines()
            rows = list(csv.reader(content, delimiter=delimiter))

            headers = rows[0] if rows else []
            data = rows[1:] if len(rows) > 1 else []

            tables.append({
                "headers": headers,
                "rows": data,
                "row_count": len(data),
                "col_count": len(headers),
                "filename": filename
            })

        elif ext == "xlsx" and openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(file["bytes"]), read_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))

            headers = rows[0] if rows else []
            data = rows[1:] if len(rows) > 1 else []

            tables.append({
                "headers": list(headers),
                "rows": data,
                "row_count": len(data),
                "col_count": len(headers),
                "filename": filename
            })

        text_blocks.append({
            "content": f"Tabela extraída do arquivo {filename} ({len(data)} linhas)",
            "metadata": {
                "type": "table_summary",
                "filename": filename
            }
        })

        return {
            "table_blocks": tables,
            "text_blocks": text_blocks
        }

    def _interpret_pdf(self, file):
        filename = file["filename"]
        raw = file["bytes"]

        # Preferencial: pdfplumber
        if pdfplumber:
            pages = []
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages.append(text)

            return {
                "text_blocks": [{
                    "content": "\n".join(pages),
                    "metadata": {
                        "type": "pdf",
                        "parser": "pdfplumber",
                        "filename": filename,
                        "pages": len(pages)
                    }
                }]
            }

        # Fallback: PyPDF2
        if PyPDF2:
            reader = PyPDF2.PdfReader(io.BytesIO(raw))
            pages = [(p.extract_text() or "") for p in reader.pages]

            return {
                "text_blocks": [{
                    "content": "\n".join(pages),
                    "metadata": {
                        "type": "pdf",
                        "parser": "pypdf2",
                        "filename": filename,
                        "pages": len(pages)
                    }
                }]
            }

        raise RuntimeError("Nenhum parser de PDF disponível (pdfplumber ou PyPDF2)")

    def _interpret_docx(self, file):
        if not Document:
            raise RuntimeError("python-docx não disponível")

        doc = Document(io.BytesIO(file["bytes"]))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        return {
            "text_blocks": [{
                "content": "\n".join(paragraphs),
                "metadata": {
                    "type": "docx",
                    "filename": file["filename"]
                }
            }]
        }

    def _interpret_image(self, file):
        return {
            "images": [{
                "filename": file["filename"],
                "content_type": file["content_type"],
                "bytes": file["bytes"],
                "metadata": {
                    "type": "image",
                    "status": "ignored_for_now"
                }
            }]
        }

    def _interpret_fallback(self, file):
        raw = file["bytes"].decode("utf-8", errors="ignore")

        return {
            "text_blocks": [{
                "content": raw,
                "metadata": {
                    "type": "unknown",
                    "filename": file["filename"]
                }
            }]
        }

    # =====================================================
    # Utils
    # =====================================================

    def _combine_text_blocks(self, blocks):
        output = []

        for block in blocks:
            meta = block.get("metadata", {})
            filename = meta.get("filename", "unknown")

            output.append(f"\n--- Arquivo: {filename} ---\n")
            output.append(block["content"])

        return "\n".join(output)
